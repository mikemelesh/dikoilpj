from io import BytesIO
from typing import Any, Dict, List, Mapping, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from .export_branding import (
    APP_NAME,
    APP_SUBTITLE,
    APP_TAGLINE,
    COLOR_ACCENT_LINE,
    COLOR_BRAND_ACCENT,
    COLOR_BRAND_DARK,
    COLOR_BRAND_MID,
    COLOR_META_BG,
    COLOR_SECTION_BG,
    COLOR_TEXT_MUTED,
    COLOR_WHITE,
    COLOR_ZEBRA,
    DATA_SECTION_TITLE_RU,
    DOCUMENT_FOOTER,
    RECORD_COUNT_LABEL_RU,
)
from .export_labels import (
    FILTERS_SECTION_TITLE_RU,
    REPORT_AUTHOR_LABEL_RU,
    REPORT_GENERATED_LABEL_RU,
    header_label,
)

_THIN_BORDER = Border(
    left=Side(style="thin", color="B4B4B4"),
    right=Side(style="thin", color="B4B4B4"),
    top=Side(style="thin", color="B4B4B4"),
    bottom=Side(style="thin", color="B4B4B4"),
)
_CARD_BORDER = Border(
    left=Side(style="thin", color="8EA9DB"),
    right=Side(style="thin", color="8EA9DB"),
    top=Side(style="thin", color="8EA9DB"),
    bottom=Side(style="thin", color="8EA9DB"),
)
_RIGHT_ALIGN_HEADERS = ("BYN", "Сумма", "Скидка", "Итого", "Оплачено", "Записей", "Рейтинг", "№")


def _sanitize_excel_sheet_title(title: str) -> str:
    if title is None:
        title = ""
    cleaned = str(title).strip()
    invalid_chars = {":", "\\", "/", "?", "*", "[", "]"}
    cleaned = "".join(ch for ch in cleaned if ch not in invalid_chars)
    cleaned = cleaned.strip() or "Отчёт"
    return cleaned[:31]


def _column_count(data: List[Dict[str, Any]]) -> int:
    if data:
        return max(len(data[0].keys()), 4)
    return 4


def _apply_fill(cell, color: str) -> None:
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")


def _merge_row(ws, row: int, col_count: int) -> str:
    last_col = get_column_letter(col_count)
    ws.merge_cells(f"A{row}:{last_col}{row}")
    return last_col


def _is_numeric_column(header: str) -> bool:
    return any(marker in header for marker in _RIGHT_ALIGN_HEADERS)


def _write_excel_letterhead(ws, col_count: int) -> int:
    row = 1
    _merge_row(ws, row, col_count)
    name_cell = ws.cell(row=row, column=1, value=APP_NAME.upper())
    name_cell.font = Font(bold=True, size=20, color=COLOR_WHITE, name="Calibri")
    name_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    _apply_fill(name_cell, COLOR_BRAND_DARK)
    ws.row_dimensions[row].height = 36

    row = 2
    mid = max(col_count // 2, 1)
    if mid > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=mid)
    if col_count > mid:
        ws.merge_cells(start_row=row, start_column=mid + 1, end_row=row, end_column=col_count)

    sub_cell = ws.cell(row=row, column=1, value=APP_SUBTITLE)
    sub_cell.font = Font(size=10, color=COLOR_TEXT_MUTED, name="Calibri")
    sub_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    _apply_fill(sub_cell, COLOR_META_BG)

    tag_cell = ws.cell(row=row, column=mid + 1, value=APP_TAGLINE)
    tag_cell.font = Font(size=9, italic=True, color=COLOR_BRAND_MID, name="Calibri")
    tag_cell.alignment = Alignment(horizontal="right", vertical="center")
    _apply_fill(tag_cell, COLOR_META_BG)
    ws.row_dimensions[row].height = 22

    row = 3
    _merge_row(ws, row, col_count)
    line_cell = ws.cell(row=row, column=1, value="")
    line_cell.border = Border(bottom=Side(style="medium", color=COLOR_ACCENT_LINE))
    ws.row_dimensions[row].height = 6

    return row + 2


def _write_excel_report_title(ws, title: str, col_count: int, start_row: int) -> int:
    row = start_row
    _merge_row(ws, row, col_count)
    cell = ws.cell(row=row, column=1, value=title.upper())
    cell.font = Font(bold=True, size=14, color=COLOR_BRAND_MID, name="Calibri")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 28
    return row + 2


def _write_excel_info_card(
    ws,
    report_meta: Dict[str, Any],
    record_count: int,
    col_count: int,
    start_row: int,
) -> int:
    items: List[tuple[str, str]] = []
    author = report_meta.get(REPORT_AUTHOR_LABEL_RU)
    generated_at = report_meta.get(REPORT_GENERATED_LABEL_RU)
    if author:
        items.append((REPORT_AUTHOR_LABEL_RU, author))
    if generated_at:
        items.append((REPORT_GENERATED_LABEL_RU, generated_at))
    items.append((RECORD_COUNT_LABEL_RU, str(record_count)))

    if not items:
        return start_row

    row = start_row
    mid = max(col_count // 2, 2)
    pairs_per_row = 2 if col_count >= 4 else 1

    for i in range(0, len(items), pairs_per_row):
        chunk = items[i : i + pairs_per_row]
        for j, (label, value) in enumerate(chunk):
            label_col = 1 + j * mid if pairs_per_row == 2 else 1
            value_col = label_col + 1
            if pairs_per_row == 2 and j == 1:
                label_col = mid + 1
                value_col = col_count

            label_cell = ws.cell(row=row, column=label_col, value=f"{label}:")
            label_cell.font = Font(bold=True, size=10, color=COLOR_TEXT_MUTED, name="Calibri")
            _apply_fill(label_cell, COLOR_META_BG)
            label_cell.border = _CARD_BORDER
            label_cell.alignment = Alignment(horizontal="right", vertical="center")

            value_cell = ws.cell(row=row, column=value_col, value=value)
            value_cell.font = Font(size=10, name="Calibri")
            value_cell.border = _CARD_BORDER
            value_cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row].height = 22
        row += 1

    return row + 1


def _write_excel_filters(ws, filters: Dict[str, Any], col_count: int, start_row: int) -> int:
    if not filters:
        return start_row

    row = start_row
    _merge_row(ws, row, col_count)
    header_cell = ws.cell(row=row, column=1, value=FILTERS_SECTION_TITLE_RU)
    header_cell.font = Font(bold=True, size=11, color=COLOR_BRAND_DARK, name="Calibri")
    header_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    _apply_fill(header_cell, COLOR_SECTION_BG)
    ws.row_dimensions[row].height = 24
    row += 1

    value_span = max(col_count - 1, 1)
    for label, value in filters.items():
        param_cell = ws.cell(row=row, column=1, value=label)
        param_cell.font = Font(bold=True, size=10, color=COLOR_TEXT_MUTED, name="Calibri")
        _apply_fill(param_cell, COLOR_META_BG)
        param_cell.border = _THIN_BORDER
        param_cell.alignment = Alignment(vertical="center")

        if value_span > 1:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=value_span + 1)
        value_cell = ws.cell(row=row, column=2, value=value)
        value_cell.font = Font(size=10, name="Calibri")
        value_cell.border = _THIN_BORDER
        value_cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 20
        row += 1

    return row + 1


def _write_excel_data_section_header(ws, col_count: int, start_row: int) -> int:
    row = start_row
    _merge_row(ws, row, col_count)
    cell = ws.cell(row=row, column=1, value=DATA_SECTION_TITLE_RU)
    cell.font = Font(bold=True, size=11, color=COLOR_BRAND_DARK, name="Calibri")
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    _apply_fill(cell, COLOR_SECTION_BG)
    ws.row_dimensions[row].height = 24
    return row + 1


def _write_excel_table(
    ws,
    data: List[Dict[str, Any]],
    start_row: int,
    column_labels: Optional[Mapping[str, str]] = None,
) -> int:
    if not data:
        cell = ws.cell(row=start_row, column=1, value="Нет данных для отображения")
        cell.font = Font(bold=True, italic=True, color="7F7F7F", name="Calibri")
        return start_row + 2

    headers = list(data[0].keys())
    row = start_row

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = (
            header_label(header, column_labels)
            if header in (column_labels or {})
            else header
        )
        cell.font = Font(bold=True, color=COLOR_WHITE, name="Calibri", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        _apply_fill(cell, COLOR_BRAND_ACCENT)
        cell.border = _THIN_BORDER

    ws.row_dimensions[row].height = 26
    data_start_row = row
    data_row_index = 0

    for item in data:
        row += 1
        data_row_index += 1
        zebra = data_row_index % 2 == 0
        for col, key in enumerate(headers, 1):
            value = item.get(key, "")
            if isinstance(value, dict):
                value = str(value)
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = Font(size=10, name="Calibri")
            cell.border = _THIN_BORDER
            if zebra:
                _apply_fill(cell, COLOR_ZEBRA)
            if _is_numeric_column(key):
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = ws.cell(row=data_start_row + 1, column=1)
    return row + 2


def _write_excel_footer(ws, col_count: int, start_row: int) -> None:
    row = start_row
    _merge_row(ws, row, col_count)
    cell = ws.cell(row=row, column=1, value=DOCUMENT_FOOTER)
    cell.font = Font(size=8, italic=True, color=COLOR_TEXT_MUTED, name="Calibri")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(top=Side(style="thin", color="B4B4B4"))
    ws.row_dimensions[row].height = 28


def _autosize_excel_columns(ws, min_row: int = 4) -> None:
    max_col = ws.max_column
    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        max_length = 0
        for row_cells in ws.iter_rows(
            min_row=min_row,
            min_col=col_idx,
            max_col=col_idx,
            values_only=False,
        ):
            cell = row_cells[0]
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))
        adjusted_width = min(max(max_length + 2, 10), 48)
        ws.column_dimensions[letter].width = adjusted_width


def generate_excel(
    data: List[Dict[str, Any]],
    title: str,
    filters: Dict[str, Any] = None,
    report_meta: Optional[Dict[str, Any]] = None,
    column_labels: Optional[Mapping[str, str]] = None,
) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = _sanitize_excel_sheet_title(title)
    ws.sheet_view.showGridLines = False

    col_count = _column_count(data)
    row = _write_excel_letterhead(ws, col_count)
    row = _write_excel_report_title(ws, title, col_count, row)
    row = _write_excel_info_card(ws, report_meta or {}, len(data), col_count, row)
    row = _write_excel_filters(ws, filters or {}, col_count, row)
    row = _write_excel_data_section_header(ws, col_count, row)
    row = _write_excel_table(ws, data, row, column_labels)
    _write_excel_footer(ws, col_count, row)
    _autosize_excel_columns(ws)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ---------------------------------------------------------------------------
# Word (DOCX)
# ---------------------------------------------------------------------------


def _hex_to_rgb(hex_color: str) -> RGBColor:
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def _set_cell_shading(cell, hex_color: str) -> None:
    shading = cell._tc.get_or_add_tcPr()
    shading.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))


def _style_table_header_cell(cell) -> None:
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].runs[0].font.color.rgb = _hex_to_rgb(COLOR_WHITE)
    cell.paragraphs[0].runs[0].font.size = Pt(9)
    _set_cell_shading(cell, COLOR_BRAND_ACCENT)


def _add_docx_horizontal_line(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    p_pr.append(
        parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="{COLOR_ACCENT_LINE}"/>'
            f"</w:pBdr>"
        )
    )


def _add_docx_letterhead(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells

    name_p = left.paragraphs[0]
    name_run = name_p.add_run(APP_NAME.upper())
    name_run.bold = True
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = _hex_to_rgb(COLOR_BRAND_DARK)

    sub_p = left.add_paragraph(APP_SUBTITLE)
    sub_p.runs[0].font.size = Pt(9)
    sub_p.runs[0].font.color.rgb = _hex_to_rgb(COLOR_TEXT_MUTED)

    tag_p = right.paragraphs[0]
    tag_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tag_run = tag_p.add_run(APP_TAGLINE)
    tag_run.italic = True
    tag_run.font.size = Pt(9)
    tag_run.font.color.rgb = _hex_to_rgb(COLOR_BRAND_MID)

    _set_cell_shading(left, COLOR_META_BG)
    _set_cell_shading(right, COLOR_META_BG)
    _add_docx_horizontal_line(doc)


def _add_docx_report_title(doc: Document, title: str) -> None:
    p = doc.add_paragraph(title.upper())
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = _hex_to_rgb(COLOR_BRAND_MID)
    p.paragraph_format.space_after = Pt(10)


def _add_docx_info_card(doc: Document, report_meta: Dict[str, Any], record_count: int) -> None:
    items: List[tuple[str, str]] = []
    author = report_meta.get(REPORT_AUTHOR_LABEL_RU)
    generated_at = report_meta.get(REPORT_GENERATED_LABEL_RU)
    if author:
        items.append((REPORT_AUTHOR_LABEL_RU, author))
    if generated_at:
        items.append((REPORT_GENERATED_LABEL_RU, generated_at))
    items.append((RECORD_COUNT_LABEL_RU, str(record_count)))
    if not items:
        return

    table = doc.add_table(rows=len(items), cols=2, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (label, value) in enumerate(items):
        label_cell, value_cell = table.rows[i].cells
        label_cell.text = f"{label}:"
        value_cell.text = value
        label_cell.paragraphs[0].runs[0].bold = True
        label_cell.paragraphs[0].runs[0].font.size = Pt(9)
        label_cell.paragraphs[0].runs[0].font.color.rgb = _hex_to_rgb(COLOR_TEXT_MUTED)
        value_cell.paragraphs[0].runs[0].font.size = Pt(9)
        _set_cell_shading(label_cell, COLOR_META_BG)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def _add_docx_filters(doc: Document, filters: Dict[str, Any]) -> None:
    if not filters:
        return

    heading = doc.add_paragraph(FILTERS_SECTION_TITLE_RU)
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(11)
    heading.runs[0].font.color.rgb = _hex_to_rgb(COLOR_BRAND_DARK)

    table = doc.add_table(rows=1, cols=2, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    hdr[0].text = "Параметр"
    hdr[1].text = "Значение"
    for cell in hdr:
        _style_table_header_cell(cell)

    for label, value in filters.items():
        row_cells = table.add_row().cells
        row_cells[0].text = str(label)
        row_cells[1].text = str(value)
        row_cells[0].paragraphs[0].runs[0].bold = True
        row_cells[0].paragraphs[0].runs[0].font.size = Pt(9)
        _set_cell_shading(row_cells[0], COLOR_META_BG)
        row_cells[1].paragraphs[0].runs[0].font.size = Pt(9)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def _add_docx_data_section(doc: Document, data: List[Dict[str, Any]], column_labels: Optional[Mapping[str, str]]) -> None:
    heading = doc.add_paragraph(DATA_SECTION_TITLE_RU)
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(11)
    heading.runs[0].font.color.rgb = _hex_to_rgb(COLOR_BRAND_DARK)

    if not data:
        p = doc.add_paragraph("Нет данных для отображения.")
        p.runs[0].italic = True
        p.runs[0].font.color.rgb = RGBColor(127, 127, 127)
        return

    table = doc.add_table(rows=1, cols=len(data[0]), style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = list(data[0].keys())

    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = (
            header_label(header, column_labels)
            if header in (column_labels or {})
            else header
        )
        _style_table_header_cell(hdr_cells[i])

    for row_idx, item in enumerate(data):
        row_cells = table.add_row().cells
        for i, key in enumerate(headers):
            value = item.get(key, "")
            if isinstance(value, dict):
                value = str(value)
            row_cells[i].text = str(value)
            row_cells[i].paragraphs[0].runs[0].font.size = Pt(9)
            if row_idx % 2 == 1:
                _set_cell_shading(row_cells[i], COLOR_ZEBRA)


def _add_docx_footer(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p_pr = p._p.get_or_add_pPr()
    p_pr.append(
        parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="4" w:color="B4B4B4"/>'
            f"</w:pBdr>"
        )
    )
    footer = doc.add_paragraph(DOCUMENT_FOOTER)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].italic = True
    footer.runs[0].font.color.rgb = _hex_to_rgb(COLOR_TEXT_MUTED)


def generate_docx(
    data: List[Dict[str, Any]],
    title: str,
    filters: Dict[str, Any] = None,
    report_meta: Optional[Dict[str, Any]] = None,
    column_labels: Optional[Mapping[str, str]] = None,
) -> BytesIO:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    _add_docx_letterhead(doc)
    _add_docx_report_title(doc, title)
    _add_docx_info_card(doc, report_meta or {}, len(data))
    _add_docx_filters(doc, filters or {})
    _add_docx_data_section(doc, data, column_labels)
    _add_docx_footer(doc)

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
